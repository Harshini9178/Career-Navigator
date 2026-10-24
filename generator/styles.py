import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_document_margins(section, top=1.0, bottom=1.0, left=1.25, right=1.0):
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.page_width = Inches(8.27)  # A4 standard
    section.page_height = Inches(11.69)

def configure_styles(doc):
    # Set default Normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    
    p_format = style.paragraph_format
    p_format.line_spacing = 1.5
    p_format.space_after = Pt(6)
    p_format.space_before = Pt(0)
    p_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, border_hex="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_page_number_field(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def setup_header_footer(section, is_prelims=False):
    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if not is_prelims:
        hrun = hp.add_run("CAREER NAVIGATOR: AI-POWERED CAREER AND PLACEMENT PLATFORM")
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # We create a 1-row, 3-column borderless table for footer
    tbl = footer.add_table(1, 3, Inches(6.00))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tbl._tbl.tblPr.append(tblBorders)
    
    # Col 1: Left
    cell_left = tbl.cell(0, 0)
    cell_left.width = Inches(2.7)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_left.paragraph_format.line_spacing = 1.0
    p_left.paragraph_format.space_after = Pt(0)
    r_left = p_left.add_run("Department of CS&E, KIT, Tiptur")
    r_left.font.name = 'Times New Roman'
    r_left.font.size = Pt(9.0)
    r_left.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Col 2: Middle
    cell_mid = tbl.cell(0, 1)
    cell_mid.width = Inches(1.3)
    p_mid = cell_mid.paragraphs[0]
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.paragraph_format.line_spacing = 1.0
    p_mid.paragraph_format.space_after = Pt(0)
    r_mid = p_mid.add_run("2026-2027")
    r_mid.font.name = 'Times New Roman'
    r_mid.font.size = Pt(9.0)
    r_mid.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Col 3: Right
    cell_right = tbl.cell(0, 2)
    cell_right.width = Inches(2.0)
    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.line_spacing = 1.0
    p_right.paragraph_format.space_after = Pt(0)
    r_right_text = p_right.add_run("Page ")
    r_right_text.font.name = 'Times New Roman'
    r_right_text.font.size = Pt(9.0)
    r_right_text.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    r_page = p_right.add_run()
    r_page.font.name = 'Times New Roman'
    r_page.font.size = Pt(9.0)
    r_page.font.bold = True
    r_page.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    add_page_number_field(r_page)

def add_chapter_heading(doc, chapter_num_text, chapter_title_text):
    # Chapter number left-justified (size 16), Title in capital & center justified (size 16-18 Bold)
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_num.paragraph_format.space_before = Pt(14)
    p_num.paragraph_format.space_after = Pt(4)
    p_num.paragraph_format.keep_with_next = True
    r_num = p_num.add_run(chapter_num_text.upper())
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(16)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(16)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(chapter_title_text.upper())
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

def add_section_heading(doc, sec_num, sec_title):
    # Section headings: 14 pt, bold, left justified
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{sec_num}  {sec_title}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

def add_subsection_heading(doc, subsec_num, subsec_title):
    # Subsection headings: 12 pt, bold, left justified
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{subsec_num}  {subsec_title}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def add_subsubsection_heading(doc, subsubsec_num, subsubsec_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{subsubsec_num}  {subsubsec_title}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_paragraph(doc, text, bold_prefix=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_bullet(doc, text, bold_prefix=None):
    # Ensure exactly ONE bullet marker is used by stripping accidental literal bullets from prefix
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    
    if bold_prefix:
        clean_prefix = bold_prefix.lstrip('•·-\t ').strip()
        if clean_prefix:
            r_pre = p.add_run(clean_prefix + " ")
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_numbered_item(doc, num_str, text, bold_prefix=None):
    # Clean hanging indent for numbered lists
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    
    r_num = p.add_run(num_str + "\t")
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(12)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    if bold_prefix:
        clean_prefix = bold_prefix.lstrip('•·-\t ').strip()
        if clean_prefix:
            r_pre = p.add_run(clean_prefix + " ")
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_table_caption(doc, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(caption_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

def add_figure(doc, img_path, caption_text, width_inches=5.6):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

def add_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Repeat header row on every page
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    # Format Header Row
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        set_cell_background(cell, "1E3A8A")  # Deep Navy
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(header_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Format Data Rows
    for row_idx, row_values in enumerate(rows_data):
        bg = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        trPr = table.rows[row_idx + 1]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        for col_idx, val in enumerate(row_values):
            cell = table.cell(row_idx + 1, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            
    # Set Column Widths if provided
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_code_box(doc, code_snippet, language="TypeScript"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    # Left border accent
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_after = Pt(0)
    
    r_tag = p.add_run(f"// --- Code Listing: {language} ---\n")
    r_tag.font.name = 'Consolas'
    r_tag.font.size = Pt(8.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    r_code = p.add_run(code_snippet)
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_callout_box(doc, title, text_content):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="1D4ED8"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(0)
    
    r_title = p.add_run(f"{title}: ")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    
    r_text = p.add_run(text_content)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(11)
    r_text.font.italic = True
    r_text.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
